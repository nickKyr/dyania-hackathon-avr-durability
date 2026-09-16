import json
import pathlib
import sys
import pandas as pd
from openpyxl import load_workbook

JSON_DIR = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path("data/llm_json")
OUT = "data/structured_extraction.xlsx"
MODEL_NAME = "claude-sonnet (subagent pass, 16 Sep 2026)"


def num(v):
    try:
        return float(v) if v is not None and v != "" else None
    except (TypeError, ValueError):
        return None


def load():
    docs, bad = [], []
    for f in sorted(JSON_DIR.glob("N*.json")):
        try:
            docs.append(json.loads(f.read_text()))
        except json.JSONDecodeError as e:
            bad.append((f.name, str(e)))
    return docs, bad


def flatten(docs):
    notes, ops, status, echo, statements, reint, ctx = [], [], [], [], [], [], []
    for d in docs:
        nid, pat, yr = d.get("note_id"), d.get("patient"), d.get("service_year")
        base = dict(note_id=nid, patient=pat, service_year=yr, note_type=d.get("note_type"), method=MODEL_NAME)
        o = d.get("operation_in_this_note") or {}
        ps = d.get("prosthetic_aortic_valve_status") or {}
        c = d.get("clinical_context") or {}
        com = c.get("comorbidities") or {}
        notes.append(dict(base, is_operative_report=d.get("is_operative_report"), n_echo_studies=len(d.get("echo_studies") or []), n_status_statements=len(d.get("prosthesis_status_statements") or []), n_reinterventions=len(d.get("reinterventions_on_aortic_valve") or []), dates_in_text="; ".join(d.get("dates_in_text") or []), extraction_notes=d.get("extraction_notes")))
        ops.append(dict(base, procedure_summary=o.get("procedure_summary"), aortic_valve_intervention=o.get("aortic_valve_intervention"), valve_model_as_written=o.get("valve_model_as_written"), valve_model_normalised=o.get("valve_model_normalised"), manufacturer=o.get("manufacturer"), valve_size_mm=num(o.get("valve_size_mm")), tavr_access=o.get("tavr_access"), concomitant_procedures="; ".join(o.get("concomitant_procedures") or []), native_valve_morphology=o.get("native_valve_morphology"), prior_cardiac_surgery_mentioned=o.get("prior_cardiac_surgery_mentioned"), cross_clamp_minutes=num(o.get("cross_clamp_minutes")), bypass_minutes=num(o.get("bypass_minutes")), evidence=o.get("evidence")))
        status.append(dict(base, has_prosthetic_aortic_valve=ps.get("patient_has_prosthetic_aortic_valve_at_time_of_note"), approach_of_existing_valve=ps.get("approach_of_existing_valve"), existing_valve_model_as_written=ps.get("existing_valve_model_as_written"), existing_valve_model_normalised=ps.get("existing_valve_model_normalised"), existing_valve_size_mm=num(ps.get("existing_valve_size_mm")), implant_date_as_written=ps.get("implant_date_as_written"), evidence=ps.get("evidence")))
        for i, e in enumerate(d.get("echo_studies") or []):
            echo.append(dict(base, study_index=i + 1, study_label=e.get("study_label"), date_as_written=e.get("date_as_written"), timing=e.get("timing_relative_to_aortic_valve_intervention"), valve_assessed=e.get("valve_assessed"), lvef_percent=num(e.get("lvef_percent")), peak_gradient_mmhg=num(e.get("peak_gradient_mmhg")), mean_gradient_mmhg=num(e.get("mean_gradient_mmhg")), dvi=num(e.get("dvi")), aortic_valve_area_cm2=num(e.get("aortic_valve_area_cm2")), aortic_valve_area_indexed_cm2_m2=num(e.get("aortic_valve_area_indexed_cm2_m2")), peak_velocity_m_s=num(e.get("peak_velocity_m_s")), lvot_velocity_or_vti=e.get("lvot_velocity_or_vti"), aortic_regurgitation_grade=e.get("aortic_regurgitation_grade"), regurgitation_location=e.get("regurgitation_location"), stenosis_severity_as_written=e.get("aortic_stenosis_severity_as_written"), leaflet_or_prosthesis_description=e.get("leaflet_or_prosthesis_description"), other_valves_and_ventricle=e.get("other_valves_and_ventricle"), evidence=e.get("evidence")))
        for s in d.get("prosthesis_status_statements") or []:
            statements.append(dict(base, statement=s.get("statement"), category=s.get("category"), date_as_written=s.get("date_as_written")))
        for r in d.get("reinterventions_on_aortic_valve") or []:
            reint.append(dict(base, type=r.get("type"), date_as_written=r.get("date_as_written"), reason_as_written=r.get("reason_as_written"), device_as_written=r.get("device_as_written"), evidence=r.get("evidence")))
        ctx.append(dict(base, sex_as_written=c.get("sex_as_written"), age_as_written=c.get("age_as_written"), height_weight_bsa_bmi=c.get("height_weight_bsa_bmi_as_written"), nyha_or_symptoms=c.get("nyha_or_symptoms"), **{f"comorbidity_{k}": v for k, v in com.items()}, antithrombotic_therapy=c.get("antithrombotic_therapy_as_written"), statin=c.get("statin_as_written")))
    return {k: pd.DataFrame(v) for k, v in dict(llm_notes=notes, llm_operations=ops, llm_prosthesis_status=status, llm_echo_studies=echo, llm_status_statements=statements, llm_reinterventions=reint, llm_clinical_context=ctx).items()}


def most_specific(names):
    names = [n for n in names if n]
    if not names:
        return None
    return max(names, key=lambda n: (len(n.split()), len(n)))


def patient_view(frames, regex_patients):
    echo = frames["llm_echo_studies"]
    ops = frames["llm_operations"]
    st = frames["llm_prosthesis_status"]
    ri = frames["llm_reinterventions"]
    ctx = frames["llm_clinical_context"]
    rows = []
    for p in sorted(set(echo.patient) | set(ops.patient) | set(st.patient)):
        o = ops[(ops.patient == p) & ops.aortic_valve_intervention.isin(["SAVR", "TAVR", "valve-in-valve TAVR", "redo SAVR"])].sort_values("service_year")
        s = st[(st.patient == p) & (st.has_prosthetic_aortic_valve == "yes")]
        pe = echo[(echo.patient == p) & (echo.valve_assessed == "prosthetic")].sort_values(["service_year", "study_index"])
        ne = echo[(echo.patient == p) & (echo.valve_assessed == "native")].sort_values(["service_year", "study_index"])
        c = ctx[ctx.patient == p]
        def series(df, col):
            return "; ".join(f"{int(y)}{'/' + str(d) if d else ''}: {v:g}" for y, d, v in zip(df.service_year, df.date_as_written.fillna(""), df[col]) if pd.notna(v))
        rows.append(dict(
            patient=p,
            llm_index_intervention=o.aortic_valve_intervention.iloc[0] if len(o) else (s.approach_of_existing_valve.iloc[0] if len(s) else None),
            llm_index_year=int(o.service_year.iloc[0]) if len(o) else None,
            llm_all_interventions="; ".join(f"{int(y)}: {t}" for y, t in zip(o.service_year, o.aortic_valve_intervention)),
            llm_valve_model=most_specific(list(o.valve_model_normalised.dropna()) + list(st[st.patient == p].existing_valve_model_normalised.dropna())),
            llm_valve_model_as_written=" | ".join(dict.fromkeys(list(o.valve_model_as_written.dropna()) + list(st[st.patient == p].existing_valve_model_as_written.dropna()))),
            llm_valve_model_variants_n=len(set(list(o.valve_model_normalised.dropna()) + list(st[st.patient == p].existing_valve_model_normalised.dropna()))),
            llm_valve_size_mm=o.valve_size_mm.dropna().iloc[0] if o.valve_size_mm.notna().any() else (s.existing_valve_size_mm.dropna().iloc[0] if s.existing_valve_size_mm.notna().any() else None),
            llm_implant_date_as_written="; ".join(s.implant_date_as_written.dropna().unique()),
            llm_n_prosthetic_echo_studies=len(pe), llm_n_native_echo_studies=len(ne),
            llm_prosthetic_mean_gradient_series=series(pe, "mean_gradient_mmhg"), llm_prosthetic_peak_gradient_series=series(pe, "peak_gradient_mmhg"),
            llm_prosthetic_dvi_series=series(pe, "dvi"), llm_prosthetic_ava_series=series(pe, "aortic_valve_area_cm2"),
            llm_prosthetic_ar_series="; ".join(f"{int(y)}: {v}" for y, v in zip(pe.service_year, pe.aortic_regurgitation_grade) if v),
            llm_prosthesis_descriptions=" || ".join(pe.leaflet_or_prosthesis_description.dropna().unique()),
            llm_native_mean_gradient_series=series(ne, "mean_gradient_mmhg"), llm_native_ava_series=series(ne, "aortic_valve_area_cm2"),
            llm_lvef_series=series(echo[echo.patient == p].sort_values("service_year"), "lvef_percent"),
            llm_reinterventions="; ".join(f"{t} ({d or 'date n/s'}; {r or 'reason n/s'})" for t, d, r in zip(ri[ri.patient == p].type, ri[ri.patient == p].date_as_written, ri[ri.patient == p].reason_as_written)),
            llm_sex=c.sex_as_written.replace("not stated", pd.NA).dropna().mode().iloc[0] if c.sex_as_written.replace("not stated", pd.NA).dropna().size else None,
            **{f"llm_{k}": ("yes" if (c[f"comorbidity_{k}"] == "yes").any() else ("no" if (c[f"comorbidity_{k}"] == "no").any() else "not stated")) for k in ["diabetes", "chronic_kidney_disease_or_dialysis", "atrial_fibrillation", "hypertension", "hyperlipidemia", "coronary_artery_disease", "heart_failure", "amyloidosis", "endocarditis_history", "pacemaker_or_icd"] if f"comorbidity_{k}" in c.columns},
            llm_antithrombotic="; ".join(c.antithrombotic_therapy.dropna().unique()),
        ))
    view = pd.DataFrame(rows)
    return regex_patients[["patient", "implant_approach", "valve_model", "valve_size_mm", "implant_year", "prosthetic_mean_gradients_by_year", "sex"]].merge(view, on="patient", how="left")


def comparison(frames, regex_implants, regex_echo):
    ops = frames["llm_operations"]
    a = regex_implants[regex_implants.note_type == "Operative Report"][["note_id", "patient", "approach", "valve_model", "valve_size_mm"]].rename(columns={"approach": "regex_approach", "valve_model": "regex_model", "valve_size_mm": "regex_size"})
    b = ops[["note_id", "aortic_valve_intervention", "valve_model_normalised", "valve_size_mm"]].rename(columns={"aortic_valve_intervention": "llm_approach", "valve_model_normalised": "llm_model", "valve_size_mm": "llm_size"})
    m = a.merge(b, on="note_id", how="left")
    norm = lambda s: s.fillna("").astype(str).str.lower().str.replace(r"\(.*\)|tavr|redo|valve-in-valve", "", regex=True).str.strip()
    m["approach_agrees"] = m.regex_approach.fillna("").str.lower().str.contains("tavr") == m.llm_approach.fillna("").str.lower().str.contains("tavr")
    m["model_family_agrees"] = [bool(x) and bool(y) and (x.split()[0] in y or y.split()[0] in x) for x, y in zip(norm(m.regex_model), norm(m.llm_model))]
    m["size_agrees"] = (m.regex_size == m.llm_size) | (m.regex_size.isna() & m.llm_size.isna())
    rg = regex_echo[(regex_echo.parameter == "mean_gradient")].groupby("note_id").agg(regex_mean_gradients=("value", lambda s: "; ".join(f"{v:g}" for v in s)), regex_contexts=("valve_context", lambda s: "; ".join(s)))
    le = frames["llm_echo_studies"].groupby("note_id").agg(llm_mean_gradients=("mean_gradient_mmhg", lambda s: "; ".join(f"{v:g}" for v in s if pd.notna(v))), llm_valves=("valve_assessed", lambda s: "; ".join(str(v) for v in s)))
    g = rg.join(le, how="outer").reset_index()
    return m, g


def main():
    docs, bad = load()
    frames = flatten(docs)
    x = pd.ExcelFile(OUT)
    regex_patients, regex_implants, regex_echo = x.parse("patients"), x.parse("implants"), x.parse("echo_values")
    view = patient_view(frames, regex_patients)
    cmp_ops, cmp_echo = comparison(frames, regex_implants, regex_echo)
    book = load_workbook(OUT)
    for name in list(book.sheetnames):
        if name.startswith("llm_") or name in ("compare_operations", "compare_echo", "patients_llm_view"):
            del book[name]
    book.save(OUT)
    with pd.ExcelWriter(OUT, engine="openpyxl", mode="a", if_sheet_exists="replace") as xw:
        view.to_excel(xw, sheet_name="patients_llm_view", index=False)
        for name, df in frames.items():
            df.to_excel(xw, sheet_name=name, index=False)
        cmp_ops.to_excel(xw, sheet_name="compare_operations", index=False)
        cmp_echo.to_excel(xw, sheet_name="compare_echo", index=False)
        dictionary = x.parse("data_dictionary")
        extra = pd.DataFrame([
            ("patients_llm_view", "one row per patient", "note-derived (LLM)", "regex columns for reference, then the LLM reading: index intervention and year, most specific valve model across all the patient's notes, number of name variants, per-study series of prosthetic mean and peak gradient, DVI, AVA, regurgitation grade, LVEF (each value prefixed with note year and any date written in the note), native pre-operative values, reinterventions with reason, sex and comorbidities as written", "Claude Sonnet subagents, one note at a time, fixed question schema, quoted evidence per field"),
            ("llm_echo_studies", "one row per echo study per note", "note-derived (LLM)", "every measurement the note gives for that study, timing relative to the valve intervention, native or prosthetic, leaflet or prosthesis description, other findings, and the evidence sentence", "LLM"),
            ("llm_operations", "one row per note", "note-derived (LLM)", "what operation the note itself documents: intervention type, device as written and normalised, size, access, concomitant procedures, native morphology, evidence", "LLM"),
            ("llm_prosthesis_status", "one row per note", "note-derived (LLM)", "whether the patient already has a prosthetic aortic valve at the time of the note, its approach, model, size and implant date as written", "LLM"),
            ("llm_status_statements", "one row per quoted sentence", "note-derived (LLM)", "every sentence about prosthesis function, categorised (normal, stenosis/degeneration, regurgitation, leak, thrombosis, endocarditis, mismatch, failure, reintervention planned/performed)", "LLM"),
            ("llm_reinterventions", "one row per reintervention", "note-derived (LLM)", "type, date and reason as written, device, evidence", "LLM"),
            ("llm_clinical_context", "one row per note", "note-derived (LLM)", "sex, age token, body size, symptoms, comorbidities as yes/no/not stated, antithrombotic and statin as written", "LLM"),
            ("llm_notes", "one row per note", "note-derived (LLM)", "counts per note, dates surviving in the text, and the extractor's free-text notes on ambiguities and contradictions", "LLM"),
            ("compare_operations / compare_echo", "per note", "derived", "regex versus LLM on approach, model family, size and mean gradients; disagreements are the rows for a clinician to check first", "join"),
        ], columns=dictionary.columns)
        pd.concat([dictionary, extra], ignore_index=True).to_excel(xw, sheet_name="data_dictionary", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "B2"
            for col in ws.columns:
                width = min(60, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col[:200]) + 2))
                ws.column_dimensions[col[0].column_letter].width = width
    print("json files:", len(docs), "unreadable:", bad)
    print({k: v.shape for k, v in frames.items()})
    print("approach agreement:", round(cmp_ops.approach_agrees.mean(), 3), "model family agreement:", round(cmp_ops.model_family_agrees.mean(), 3), "size agreement:", round(cmp_ops.size_agrees.mean(), 3), "of", len(cmp_ops))
    e = frames["llm_echo_studies"]
    print("echo studies:", len(e), "prosthetic:", (e.valve_assessed == "prosthetic").sum(), "native:", (e.valve_assessed == "native").sum(), "unclear:", (e.valve_assessed == "unclear").sum())
    print("patients with >=2 prosthetic echo studies in different years:", (e[e.valve_assessed == "prosthetic"].groupby("patient").service_year.nunique() >= 2).sum())


if __name__ == "__main__":
    main()

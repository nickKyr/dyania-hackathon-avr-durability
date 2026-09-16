import pathlib
import pandas as pd
from openpyxl.styles import Font, PatternFill

from _inputs import require

RAW = pathlib.Path("data/raw_tables")
OUT = pathlib.Path("data/all_data_one_sheet.xlsx")
COLS = ["patient", "year", "source", "note_id", "note_type", "category", "item", "value", "unit", "detail", "evidence", "method", "duplicate_row"]


def load():
    return {p.stem: pd.read_parquet(p) for p in RAW.glob("*.parquet")}


def s(v):
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)


def rows_from_notes(R):
    notes = R["notes"].set_index("note_id")
    out = []
    for _, r in R["note_operations"].iterrows():
        if not (r.aortic_valve_intervention and s(r.aortic_valve_intervention) not in ("", "none")) and not s(r.valve_model_as_written):
            continue
        out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category="operation described in this note",
                        item=s(r.aortic_valve_intervention), value=s(r.valve_model_normalised) or s(r.valve_model_as_written), unit="", detail="; ".join(x for x in [f"as written: {s(r.valve_model_as_written)}" if s(r.valve_model_as_written) else "", f"size {s(r.valve_size_mm)} mm" if s(r.valve_size_mm) else "", s(r.manufacturer), s(r.tavr_access), s(r.concomitant_procedures), f"native valve {s(r.native_valve_morphology)}" if s(r.native_valve_morphology) else "", s(r.procedure_text)] if x),
                        evidence=s(r.evidence), method=r.method, duplicate_row=""))
    for _, r in R["note_prosthesis_status"].iterrows():
        if s(r.has_prosthetic_aortic_valve) != "yes":
            continue
        out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category="prosthetic valve in history",
                        item=s(r.approach), value=s(r.valve_model_normalised) or s(r.valve_model_as_written), unit="", detail="; ".join(x for x in [f"as written: {s(r.valve_model_as_written)}" if s(r.valve_model_as_written) else "", f"size {s(r.valve_size_mm)} mm" if s(r.valve_size_mm) else "", f"implant date as written: {s(r.implant_date_as_written)}" if s(r.implant_date_as_written) else ""] if x),
                        evidence=s(r.evidence), method=r.method, duplicate_row=""))
    E = R["echo_measurements"]
    for _, r in E.iterrows():
        if r.parameter == "study_recorded":
            continue
        cat = "echo study" + (f" {int(r.study_index)}" if s(r.study_index) else "") + (f" ({s(r.valve_assessed)} valve, {s(r.timing)})" if s(r.valve_assessed) or s(r.timing) else "")
        out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category=cat,
                        item=r.parameter, value=s(r.value_num) if s(r.value_num) else s(r.value_text), unit=s(r.unit), detail="; ".join(x for x in [s(r.study_label), f"date as written: {s(r.date_as_written)}" if s(r.date_as_written) else ""] if x),
                        evidence=s(r.evidence), method=r.method, duplicate_row=""))
    for _, r in R["prosthesis_statements"].iterrows():
        out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category="statement about prosthesis",
                        item=s(r.category), value=s(r.statement), unit="", detail=f"date as written: {s(r.date_as_written)}" if s(r.date_as_written) else "", evidence=s(r.statement), method=r.method, duplicate_row=""))
    for _, r in R["reinterventions"].iterrows():
        out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category="reintervention on aortic valve",
                        item=s(r.type), value=s(r.reason_as_written), unit="", detail="; ".join(x for x in [f"date as written: {s(r.date_as_written)}" if s(r.date_as_written) else "", f"device: {s(r.device_as_written)}" if s(r.device_as_written) else ""] if x), evidence=s(r.evidence), method=r.method, duplicate_row=""))
    for _, r in R["events_regex"].iterrows():
        out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category="keyword found",
                        item=s(r.event_type), value="negated in text" if r.negated else "mentioned", unit="", detail="", evidence=s(r.snippet), method=r.method, duplicate_row=""))
    for _, r in R["clinical_context"].iterrows():
        com = {k.replace("comorbidity_", ""): v for k, v in r.items() if k.startswith("comorbidity_") and s(v) not in ("", "not stated")}
        for k, v in com.items():
            out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category="comorbidity", item=k, value=s(v), unit="", detail="", evidence="", method=r.method, duplicate_row=""))
        for k, label in [("sex_as_written", "sex"), ("age_as_written", "age"), ("height_weight_bsa_bmi_as_written", "height / weight / BSA / BMI"), ("nyha_or_symptoms", "symptoms"), ("antithrombotic_therapy_as_written", "antithrombotic therapy"), ("statin_as_written", "statin")]:
            if s(r[k]) and s(r[k]) != "not stated":
                out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category="clinical context", item=label, value=s(r[k]), unit="", detail="", evidence="", method=r.method, duplicate_row=""))
    for _, r in R["dates_in_text"].iterrows():
        out.append(dict(patient=r.patient, year=notes.service_year.get(r.note_id), source="notes", note_id=r.note_id, note_type=notes.note_type.get(r.note_id), category="date left in text", item="date string", value=s(r.date_string), unit="", detail="", evidence="", method=r.method, duplicate_row=""))
    for _, r in R["notes"].iterrows():
        out.append(dict(patient=r.patient, year=r.service_year, source="notes", note_id=r.note_id, note_type=r.note_type, category="note", item="note exists", value=f"{s(r.note_type)}, {s(r.signed_status)}", unit="", detail="; ".join(x for x in [s(r.service), s(r.author_type), s(r.author_specialty), f"{int(r.n_words)} words"] if x), evidence="", method="source file", duplicate_row=""))
    return out


def rows_from_labs(R):
    L = R["labs_raw"]
    return [dict(patient=r.Patient, year=r["Result Date"], source="labs", note_id="", note_type="", category="lab result" + (f" ({s(r['Lab Type'])})" if s(r["Lab Type"]) else ""),
                 item=r["Lab Component Name"], value=s(r["String Value"]) if s(r["String Value"]) else s(r["Numeric Value"]), unit=s(r.Unit),
                 detail="; ".join(x for x in [f"flag: {s(r.Flag)}" if s(r.Flag) else "", f"reference: {s(r['Reference Values'])}" if s(r["Reference Values"]) else "", f"LOINC {s(r['Loinc Code'])}" if s(r["Loinc Code"]) else ""] if x),
                 evidence="", method="source file", duplicate_row="yes" if r.is_exact_duplicate else "") for _, r in L.iterrows()]


def rows_from_meds(R):
    M = R["meds_raw"]
    out = []
    for _, r in M.iterrows():
        out.append(dict(patient=r.Patient, year=r["Start Date"], source="medications", note_id="", note_type="", category=f"medication order ({s(r.Mode).lower()})",
                        item=s(r["Simple Generic Name"]) or s(r["Proper Name"]), value=" ".join(x for x in [s(r.Dose), s(r["Dose Unit"]), s(r.Frequency)] if x), unit="",
                        detail="; ".join(x for x in [s(r["Proper Name"]), s(r.Route), s(r["Medication Therapeutic Class"]), f"end {int(r['End Date'])}" if s(r["End Date"]) else "", f"stopped {int(r['Discontinued Date'])}: {s(r['Discontinued Reason'])}" if s(r["Discontinued Date"]) else ""] if x),
                        evidence="", method="source file", duplicate_row="yes" if r.is_exact_duplicate else ""))
    for _, r in R["diagnoses"].iterrows():
        out.append(dict(patient=r.patient, year=r.start_year, source="medications", note_id="", note_type="", category="diagnosis linked to a medication order", item=s(r.icd10_code), value=s(r.diagnosis_name), unit="", detail=s(r.icd10_name), evidence="", method="source file", duplicate_row=""))
    return out


def main():
    require(RAW, hint="Run scripts/03_build_raw_tables.py first.")
    R = load()
    rows = rows_from_notes(R) + rows_from_labs(R) + rows_from_meds(R)
    df = pd.DataFrame(rows)[COLS]
    df["year"] = pd.to_numeric(df.year, errors="coerce")
    order = {"note": 0, "operation described in this note": 1, "prosthetic valve in history": 2, "reintervention on aortic valve": 3, "statement about prosthesis": 4, "keyword found": 5, "comorbidity": 7, "clinical context": 8, "date left in text": 9}
    df["_o"] = df.category.map(lambda c: order.get(c, 6 if c.startswith("echo") else 10))
    df = df.sort_values(["patient", "year", "source", "note_id", "_o", "category", "item"]).drop(columns="_o")
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="all_data", index=False)
        ws = xw.book["all_data"]
        widths = dict(patient=13, year=7, source=12, note_id=9, note_type=16, category=34, item=30, value=40, unit=9, detail=50, evidence=70, method=22, duplicate_row=10)
        for i, c in enumerate(COLS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths[c]
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EEF7")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    print(len(df), "rows;", df.source.value_counts().to_dict())


if __name__ == "__main__":
    main()
